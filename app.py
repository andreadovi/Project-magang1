import streamlit as st
import pandas as pd
import io
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from scheduler import generate_mixing_schedule

st.set_page_config(page_title="Mixing Scheduler", page_icon="🧪", layout="wide")
st.title("🧪 Mixing Schedule Planner")

# ─── Session State ────────────────────────────────────────────────────────────
if "master_mixer"  not in st.session_state:
    st.session_state.master_mixer  = pd.DataFrame(columns=["Mixer", "Kapasitas_kg", "Grup_Cleaning"])
if "master_produk" not in st.session_state:
    st.session_state.master_produk = pd.DataFrame(columns=["Kode_Produk", "Nama_Produk", "Grup_Cleaning", "Kg_per_CS", "Mixer_Kompatibel"])
if "filling_plan"  not in st.session_state:
    st.session_state.filling_plan  = pd.DataFrame(columns=["Kode_Produk", "Nama_Produk", "Target_CS", "Tanggal_Filling", "Shift_Filling", "Urgent"])

tab1, tab2, tab3 = st.tabs(["⚙️ Master Data", "📋 Input Planning", "📅 Jadwal Mixing"])

# ═════════════════════════════════════════════════════════════════════════════
# TAB 1 — MASTER DATA
# ═════════════════════════════════════════════════════════════════════════════
with tab1:
    st.header("Master Data")
    col1, col2 = st.columns(2)

    # ── Master Mixer ──────────────────────────────────────────────────────────
    with col1:
        st.subheader("🔧 Master Mixer")

        tmpl_mixer = pd.DataFrame({
            "Mixer":         ["Mixer A", "Mixer B", "Mixer C"],
            "Kapasitas_kg":  [500, 800, 300],
            "Grup_Cleaning": ["Grup 1", "Grup 1", "Grup 2"]
        })
        buf = io.BytesIO(); tmpl_mixer.to_excel(buf, index=False)
        st.download_button("📥 Download Template Mixer", buf.getvalue(),
                           "template_master_mixer.xlsx", use_container_width=True)

        up_mixer = st.file_uploader("Upload Master Mixer", type=["xlsx", "csv"], key="mixer_upload")
        if up_mixer:
            df = pd.read_excel(up_mixer) if up_mixer.name.endswith("xlsx") else pd.read_csv(up_mixer)
            req = {"Mixer", "Kapasitas_kg", "Grup_Cleaning"}
            if req.issubset(df.columns):
                st.session_state.master_mixer = df
                st.success(f"✅ {len(df)} mixer berhasil diupload!")
            else:
                st.error(f"❌ Kolom harus: {req}")

        if not st.session_state.master_mixer.empty:
            st.dataframe(st.session_state.master_mixer, use_container_width=True, hide_index=True)

    # ── Master Produk ─────────────────────────────────────────────────────────
    with col2:
        st.subheader("📦 Master Produk")
        st.caption("**Mixer_Kompatibel**: pisahkan dengan koma. Contoh: `Mixer A, Mixer B`")

        tmpl_produk = pd.DataFrame({
            "Kode_Produk":      ["P001", "P002", "P003"],
            "Nama_Produk":      ["Produk Alpha", "Produk Beta", "Produk Gamma"],
            "Grup_Cleaning":    ["Grup 1", "Grup 1", "Grup 2"],
            "Kg_per_CS":        [12.5, 8.0, 10.0],
            "Mixer_Kompatibel": ["Mixer A, Mixer B", "Mixer B", "Mixer C"]
        })
        buf2 = io.BytesIO(); tmpl_produk.to_excel(buf2, index=False)
        st.download_button("📥 Download Template Produk", buf2.getvalue(),
                           "template_master_produk.xlsx", use_container_width=True)

        up_produk = st.file_uploader("Upload Master Produk", type=["xlsx", "csv"], key="produk_upload")
        if up_produk:
            df = pd.read_excel(up_produk) if up_produk.name.endswith("xlsx") else pd.read_csv(up_produk)
            req = {"Kode_Produk", "Nama_Produk", "Grup_Cleaning", "Kg_per_CS", "Mixer_Kompatibel"}
            if req.issubset(df.columns):
                st.session_state.master_produk = df
                st.success(f"✅ {len(df)} produk berhasil diupload!")
            else:
                st.error(f"❌ Kolom harus: {req}")

        if not st.session_state.master_produk.empty:
            st.dataframe(st.session_state.master_produk, use_container_width=True, hide_index=True)

# ═════════════════════════════════════════════════════════════════════════════
# TAB 2 — INPUT PLANNING
# ═════════════════════════════════════════════════════════════════════════════
with tab2:
    st.header("Input Planning Filling")

    if st.session_state.master_produk.empty:
        st.warning("⚠️ Upload Master Produk dulu di tab Master Data.")
    else:
        produk_options = (st.session_state.master_produk["Kode_Produk"] + " - " +
                          st.session_state.master_produk["Nama_Produk"])
        produk_map     = dict(zip(produk_options, st.session_state.master_produk["Kode_Produk"]))
        kg_map         = dict(zip(
            st.session_state.master_produk["Kode_Produk"],
            st.session_state.master_produk["Kg_per_CS"]
        ))

        with st.expander("➕ Tambah Item Planning", expanded=True):
            c1, c2, c3 = st.columns(3)
            with c1:
                sel_produk  = st.selectbox("Produk", produk_options)
                target_cs   = st.number_input("Target (CS)", min_value=1, value=100)
                kode_sel    = produk_map[sel_produk]
                kg_preview  = target_cs * float(kg_map.get(kode_sel, 0))
                st.caption(f"≈ **{kg_preview:,.1f} kg**")
            with c2:
                tgl_filling   = st.date_input("Tanggal Filling", value=datetime.today())
                shift_filling = st.selectbox("Shift Filling", [1, 2, 3])
            with c3:
                urgent = st.radio("Status", ["Urgent", "Tidak Urgent"], horizontal=True)
                st.write("")
                st.write("")
                if st.button("➕ Tambah", use_container_width=True):
                    nama = sel_produk.split(" - ", 1)[1]
                    new_row = pd.DataFrame([{
                        "Kode_Produk":     kode_sel,
                        "Nama_Produk":     nama,
                        "Target_CS":       target_cs,
                        "Tanggal_Filling": tgl_filling.strftime("%Y-%m-%d"),
                        "Shift_Filling":   shift_filling,
                        "Urgent":          urgent
                    }])
                    st.session_state.filling_plan = pd.concat(
                        [st.session_state.filling_plan, new_row], ignore_index=True)
                    st.success("✅ Berhasil ditambahkan!")
                    st.rerun()

        # ── Upload bulk ───────────────────────────────────────────────────────
        st.subheader("📤 Atau Upload Bulk")
        tmpl_plan = pd.DataFrame({
            "Kode_Produk":     ["P001", "P002"],
            "Target_CS":       [100, 50],
            "Tanggal_Filling": ["2024-01-15", "2024-01-15"],
            "Shift_Filling":   [2, 3],
            "Urgent":          ["Urgent", "Tidak Urgent"]
        })
        buf3 = io.BytesIO(); tmpl_plan.to_excel(buf3, index=False)
        st.download_button("📥 Download Template Planning", buf3.getvalue(),
                           "template_planning.xlsx", use_container_width=True)

        up_plan = st.file_uploader("Upload Planning", type=["xlsx", "csv"], key="plan_upload")
        if up_plan:
            df = pd.read_excel(up_plan) if up_plan.name.endswith("xlsx") else pd.read_csv(up_plan)
            req = {"Kode_Produk", "Target_CS", "Tanggal_Filling", "Shift_Filling", "Urgent"}
            if req.issubset(df.columns):
                df = df.merge(st.session_state.master_produk[["Kode_Produk", "Nama_Produk"]],
                              on="Kode_Produk", how="left")
                st.session_state.filling_plan = pd.concat(
                    [st.session_state.filling_plan, df], ignore_index=True)
                st.success(f"✅ {len(df)} item berhasil diupload!")
                st.rerun()
            else:
                st.error(f"❌ Kolom harus: {req}")

        # ── Tampilkan planning ────────────────────────────────────────────────
        if not st.session_state.filling_plan.empty:
            st.subheader("📋 Daftar Planning")

            # Add kg preview column
            df_disp = st.session_state.filling_plan.copy()
            df_disp["Est_kg"] = df_disp.apply(
                lambda r: f"{float(r['Target_CS']) * float(kg_map.get(r['Kode_Produk'], 0)):,.1f}",
                axis=1
            )

            def style_urgent(row):
                if row.get("Urgent") == "Urgent":
                    return ["background-color: #fff3cd"] * len(row)
                return [""] * len(row)

            st.dataframe(df_disp.style.apply(style_urgent, axis=1),
                         use_container_width=True, hide_index=True)
            st.caption("🟡 Kuning = Urgent")

            if st.button("🗑️ Reset Semua Planning", use_container_width=True):
                st.session_state.filling_plan = pd.DataFrame(
                    columns=["Kode_Produk", "Nama_Produk", "Target_CS",
                             "Tanggal_Filling", "Shift_Filling", "Urgent"])
                st.rerun()

# ═════════════════════════════════════════════════════════════════════════════
# TAB 3 — JADWAL MIXING
# ═════════════════════════════════════════════════════════════════════════════
with tab3:
    st.header("Jadwal Mixing Otomatis")

    ready = (not st.session_state.master_mixer.empty and
             not st.session_state.master_produk.empty and
             not st.session_state.filling_plan.empty)

    if not ready:
        st.warning("⚠️ Lengkapi Master Mixer, Master Produk, dan Input Planning terlebih dahulu.")
    else:
        if st.button("⚡ Generate Jadwal Mixing", type="primary", use_container_width=True):
            with st.spinner("Menjadwalkan mixing..."):
                result = generate_mixing_schedule(
                    st.session_state.master_mixer,
                    st.session_state.master_produk,
                    st.session_state.filling_plan
                )
            st.session_state.schedule_result = result

        if "schedule_result" in st.session_state:
            result       = st.session_state.schedule_result
            schedule_df  = result["schedule"]
            warnings     = result["warnings"]
            shifted      = result["shifted"]
            unscheduled  = result["unscheduled"]

            # ── Warnings ─────────────────────────────────────────────────────
            if warnings:
                st.subheader("⚠️ Peringatan")
                for w in warnings:
                    st.warning(w)

            # ── Jadwal digeser ────────────────────────────────────────────────
            if shifted:
                st.subheader("🔀 Jadwal Filling Digeser (Tidak Urgent)")
                st.dataframe(pd.DataFrame(shifted).style.applymap(
                    lambda _: "background-color: #fff3cd"),
                    use_container_width=True, hide_index=True)

            # ── Tidak bisa dijadwalkan ────────────────────────────────────────
            if unscheduled:
                st.subheader("❌ Tidak Bisa Dijadwalkan")
                for u in unscheduled:
                    st.error(u)

            # ── Jadwal Mixing ─────────────────────────────────────────────────
            if not schedule_df.empty:
                st.subheader("📅 Jadwal Mixing")

                dates = sorted(schedule_df["Tanggal"].unique())
                sel_date = st.selectbox("Filter Tanggal", ["Semua"] + list(dates))
                disp_df  = schedule_df[schedule_df["Tanggal"] == sel_date] if sel_date != "Semua" else schedule_df

                def style_rows(row):
                    if row.get("Cleaning", False):
                        return ["background-color: #f8d7da"] * len(row)
                    return [""] * len(row)

                st.dataframe(disp_df.drop(columns=["Cleaning"], errors="ignore")
                             .style.apply(style_rows, axis=1),
                             use_container_width=True, hide_index=True)
                st.caption("🔴 Merah = shift cleaning | Kolom **Total_CS** = jumlah CS yang di-mixing")

                # ── Download Excel ────────────────────────────────────────────
                def to_excel(df):
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "Jadwal Mixing"

                    hdr_fill = PatternFill("solid", fgColor="1F4E79")
                    hdr_font = Font(bold=True, color="FFFFFF")
                    cln_fill = PatternFill("solid", fgColor="F8D7DA")
                    alt_fill = PatternFill("solid", fgColor="EBF1DE")

                    headers = list(df.columns)
                    for ci, h in enumerate(headers, 1):
                        cell = ws.cell(row=1, column=ci, value=h)
                        cell.fill = hdr_fill
                        cell.font = hdr_font
                        cell.alignment = Alignment(horizontal="center")

                    for ri, (_, row) in enumerate(df.iterrows(), 2):
                        is_clean = row.get("Cleaning", False)
                        for ci, val in enumerate(row, 1):
                            cell = ws.cell(row=ri, column=ci, value=val)
                            cell.fill = cln_fill if is_clean else (alt_fill if ri % 2 == 0 else PatternFill())
                            cell.alignment = Alignment(horizontal="center")

                    for col in ws.columns:
                        mx = max(len(str(c.value or "")) for c in col)
                        ws.column_dimensions[get_column_letter(col[0].column)].width = min(mx + 4, 40)

                    buf = io.BytesIO()
                    wb.save(buf)
                    return buf.getvalue()

                excel_data = to_excel(disp_df)
                st.download_button(
                    "📥 Download Excel", excel_data, "jadwal_mixing.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
